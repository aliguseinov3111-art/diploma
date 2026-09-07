# -*- coding: utf-8 -*-
"""
Донастройка столбца Q (07.09.2026): вернуть тег ККО клиентскому свету RAL 3020
раздела Р9, который является дословной позицией Альбома точек продаж:
  стр. 1498–1501  — трек RAL 3020 + трековый светильник RAL 3020 (аналог, разрешён ББ);
  стр. 1508–1512  — лайтбоксы потолочные круглые RAL 3020, h-114, диммируемые.
ORTUS SORES 650 RAL 7047 (1506–1507), линейные красные (1502–1505) и пульт (1513) НЕ трогаем.

Правка на уровне zip/XML (только xl/worksheets/sheet1.xml). Лог -> Q_lightbox_log.csv
"""
import csv, re, shutil, zipfile
from pathlib import Path
import openpyxl

HERE   = Path(__file__).resolve().parent
MASTER = HERE.parent.parent / "ТС ЕР исходник.xlsx"
LOG    = HERE / "Q_lightbox_log.csv"
SHEETXML = "xl/worksheets/sheet1.xml"
SHEET  = "Смета ЕР (общ) "
SI_KKO = 846  # sharedStrings индекс «ККО»

ROWS = list(range(1498, 1502)) + list(range(1508, 1513))   # 1498–1501, 1508–1512


def main():
    with zipfile.ZipFile(MASTER) as z:
        xml = z.read(SHEETXML).decode("utf-8")

    done, problems = [], []
    for rn in ROWS:
        # ячейка должна быть пустой стилизованной: <c r="Q{rn}" s="NN"/>
        m = re.search(r'<c r="Q%d"( s="\d+")?\s*/>' % rn, xml)
        if not m:
            # вдруг уже с тегом?
            m2 = re.search(r'<c r="Q%d"[^>]*>.*?</c>' % rn, xml)
            problems.append((rn, "не пустая ячейка: " + (m2.group(0)[:60] if m2 else "не найдена")))
            continue
        style = m.group(1) or ""
        xml = xml[:m.start()] + f'<c r="Q{rn}"{style} t="s"><v>{SI_KKO}</v></c>' + xml[m.end():]
        done.append(rn)

    if problems:
        print("ПРОБЛЕМЫ, правка не применена:")
        for x in problems:
            print("  ", x)
        return

    bak = MASTER.with_suffix(".xlsx.bak_preLB")
    shutil.copy2(MASTER, bak)
    tmp = MASTER.with_suffix(".xlsx.tmp")
    with zipfile.ZipFile(MASTER) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename == SHEETXML:
                data = xml.encode("utf-8")
            zi = zipfile.ZipInfo(it.filename, date_time=it.date_time)
            zi.compress_type = it.compress_type; zi.external_attr = it.external_attr
            zi.internal_attr = it.internal_attr; zi.create_system = it.create_system
            zout.writestr(zi, data)
    tmp.replace(MASTER)

    wb = openpyxl.load_workbook(MASTER, data_only=True)
    ws = wb[SHEET]
    with open(LOG, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Строка", "Стало", "Наименование"])
        for rn in done:
            w.writerow([rn, ws.cell(row=rn, column=17).value, str(ws.cell(row=rn, column=7).value)[:110]])

    import collections
    razd = ""
    c = collections.Counter()
    for r in ws.iter_rows(min_row=15, max_row=ws.max_row):
        v = r[16].value
        if v and str(v).strip() in ("ККО", "ИНФРАСТРУКТУРА", "ККО, ИНФРАСТРУКТУРА"):
            c[str(v).strip()] += 1
    print(f"применено: {len(done)}  (бэкап {bak.name})")
    print("Q итог:", dict(c), "= всего", sum(c.values()))


if __name__ == "__main__":
    main()
