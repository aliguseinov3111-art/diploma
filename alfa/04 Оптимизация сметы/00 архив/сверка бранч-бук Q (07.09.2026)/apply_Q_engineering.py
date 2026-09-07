# -*- coding: utf-8 -*-
"""
Второй проход по столбцу Q (07.09.2026): новое правило пользователя —
классификацию оставляем только для ОТДЕЛОЧНЫХ материалов (раздел Р3);
инженерные разделы Р5–Р9 очищаем, КРОМЕ позиций, чья номенклатурная линейка
названа в бранч-буке (не «или аналог»):
    RIFAR (Р6) · Globalvent/VTS/NED (Р7) · кондиционеры Haier/MDV (Р8) ·
    светильники REFLECTA — Rota/Quark/Ray/FAT Comfort/Longa/Agat/SILENT/Q45/Grata/MICA (Р9) ·
    розетки Legrand Mozaik (Р9) · смеситель Ideal Standard Ceraline BC193AA (Р5, дословно в ББ).
Проектный раздел Р1 (ТЗК) — очищается.

Правка на уровне zip/XML (переписывается только xl/worksheets/sheet1.xml).
Лог -> Q_engineering_log.csv
"""
import csv, re, shutil, zipfile
from pathlib import Path
import openpyxl

HERE   = Path(__file__).resolve().parent
MASTER = HERE.parent.parent / "ТС ЕР исходник.xlsx"
LOG    = HERE / "Q_engineering_log.csv"
SHEETXML = "xl/worksheets/sheet1.xml"
SHEET  = "Смета ЕР (общ) "
SI_TAGS = {846: "ККО", 6176: "ИНФРАСТРУКТУРА", 6178: "ККО, ИНФРАСТРУКТУРА"}

# --- что ОСТАЁТСЯ в инженерных разделах: линейка названа в ББ дословно
KEEP_ENG = re.compile(
    r"\bRIFAR\b"
    r"|Globalvent|\bVTS\b|\bNED\b"
    r"|\bRota\b|\bQuark\b|Ray\s*line|Ray\s*IN\b|FAT\s*Comfort|\bLonga\b|\bAgat\b"
    r"|Silent\s*acoustic|Q45\s*Track|Grata\s*Slim|\bMICA\s*D?\d"
    r"|Legrand[\s\S]{0,18}Mozaik"
    r"|IDEAL\s*STANDARD\s*CERALINE\s*BC193AA",
    re.I)
# кондиционер Haier/MDV — да; отдельная строка-комплектующая (разветвитель/помпа) — нет.
# «пульт», «панель» и т.п. в составе блока не исключают: это тот же кондиционер.
COND = re.compile(r"\b(Haier|MDV)\b", re.I)
COND_UNIT = re.compile(r"сплит|кондиционер|внутр|наружн|\bблок\b|VRF|канальн|кассетн", re.I)
COND_ACC  = re.compile(r"^\s*(разветвитель|помпа|дренажн|фреонопровод|трасса)", re.I)


def keep(razd, name):
    if razd == "Р3":
        return True
    if razd in ("Р5", "Р6", "Р7", "Р8", "Р9"):
        if KEEP_ENG.search(name):
            return True
        if COND.search(name) and COND_UNIT.search(name) and not COND_ACC.search(name):
            return True
        return False
    return False  # Р1 и всё прочее


def main():
    wb = openpyxl.load_workbook(MASTER, data_only=True)
    ws = wb[SHEET]
    razd = ""
    plan = []
    for r in ws.iter_rows(min_row=15, max_row=ws.max_row):
        g = str(r[6].value).strip() if r[6].value is not None else ""
        if re.match(r"^Раздел\s*№", g):
            m = re.match(r"^Раздел\s*№\s*(\d+)", g)
            razd = "Р" + m.group(1) if m else razd
            continue
        q = r[16].value
        if not q or str(q).strip() not in SI_TAGS.values():
            continue
        name = g
        if not keep(razd, name):
            plan.append((r[0].row, str(q).strip(), razd, name[:110]))

    print(f"К очистке: {len(plan)}")
    import collections
    print("  по разделам:", dict(collections.Counter(p[2] for p in plan)))

    with zipfile.ZipFile(MASTER) as z:
        xml = z.read(SHEETXML).decode("utf-8")

    done, problems = [], []
    for rn, oldtag, rz, name in plan:
        pat = re.compile(r'<c r="Q%d"( s="(\d+)")? t="s"><v>(\d+)</v></c>' % rn)
        m = pat.search(xml)
        if not m:
            problems.append((rn, "ячейка Q с t=s не найдена")); continue
        if SI_TAGS.get(int(m.group(3))) != oldtag:
            problems.append((rn, f"тег в файле не совпал: idx {m.group(3)}")); continue
        style = m.group(1) or ""
        xml = xml[:m.start()] + f'<c r="Q{rn}"{style}/>' + xml[m.end():]
        done.append((rn, oldtag, rz, name))

    if problems:
        print("ПРОБЛЕМЫ, правка не применена:")
        for x in problems[:20]:
            print("  ", x)
        return

    bak = MASTER.with_suffix(".xlsx.bak_preEng")
    shutil.copy2(MASTER, bak)
    tmp = MASTER.with_suffix(".xlsx.tmp")
    with zipfile.ZipFile(MASTER) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename == SHEETXML:
                data = xml.encode("utf-8")
            zi = zipfile.ZipInfo(it.filename, date_time=it.date_time)
            zi.compress_type = it.compress_type; zi.external_attr = it.external_attr
            zi.internal_attr = it.internal_attr; zi.create_system = it.create_system
            zout.writestr(zi, data)
    tmp.replace(MASTER)

    with open(LOG, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Строка", "Был тег", "Раздел", "Наименование", "Действие"])
        for rn, t, rz, name in done:
            w.writerow([rn, t, rz, name, "очищено (нет в ББ)"])

    wb2 = openpyxl.load_workbook(MASTER, data_only=True)
    ws2 = wb2[SHEET]
    cnt = collections.Counter(str(r[16].value).strip() for r in ws2.iter_rows(min_row=15, max_row=ws2.max_row)
                              if r[16].value and str(r[16].value).strip() in SI_TAGS.values())
    print(f"применено: {len(done)}   бэкап: {bak.name}")
    print("Q после второго прохода:", dict(cnt))
    print("лог:", LOG.name)


if __name__ == "__main__":
    main()
