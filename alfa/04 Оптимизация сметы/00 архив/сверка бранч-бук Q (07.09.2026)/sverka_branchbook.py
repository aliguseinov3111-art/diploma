# -*- coding: utf-8 -*-
"""
Построчная сверка столбца Q («Бранч-бук (ККО / Инфраструктура)») листа
«Смета ЕР (общ)» в  ТС ЕР исходник.xlsx  с двумя эталонными таблицами:
  00 архив/бранч бук ККО таблица.xlsx          — Альбом точек продаж (ТП), раздел 6
  00 архив/Бранч-бук Инфраструктура таблица.xlsx — Бранч-бук инфраструктуры (18 альбомов)

Результат: «Сверка Q — построчно.xlsx» в этой же папке.
Метод — сопоставление названия/примечания работы с номинированной номенклатурой
(модели, материалы, зоны), характерной для каждого альбома.
"""
import re, openpyxl
from pathlib import Path
from openpyxl.styles import Font, PatternFill, Alignment

BASE = Path(__file__).resolve().parent
SRC  = BASE.parent.parent / "ТС ЕР исходник.xlsx"   # мастер лежит на два уровня выше
OUT  = BASE / "Сверка Q — построчно.xlsx"
# ВНИМАНИЕ: Сверка Q — построчно.xlsx фиксирует состояние ПРОХОДА 1 (до чисток
# инженерки 07.09.2026). Перезапуск перезапишет её текущим (более поздним) состоянием.
SHEET = "Смета ЕР (общ) "

# ---------------------------------------------------------------- сигнатуры
# «Сильные» признаки — номинированные позиции, встречающиеся ТОЛЬКО в одном альбоме.
KKO = {
    "терраццо / Kerama SG632400R (пол ТП)":      r"террацц|sg632400",
    "плинтус Fezard 80 мм (ТП)":                  r"\bfezard\b",
    "стеклообои «паутинка» (ТП, стены)":          r"стеклообо|паутинк",
    "перегородка NAYADA-Crystal (ТП)":            r"nayada",
    "свет ТП (SILED/лайтбокс/INI LED/RINGO)":     r"siled|la\s*linea|лайтбокс|\bini\s*led\b|ringo\s*m|босма|tunic\s*led",
    "клиентский трек/шинопровод RAL 3020 (ТП)":   r"(шинопровод|\bтрек\b|трековый)[^\n]{0,60}ral\s*?30?20|ral\s*?30?20[^\n]{0,60}(шинопровод|\bтрек\b|трековый)",
    "мебель ТП (органайзер/Кракен/смартбокс)":    r"шкаф[-\s]?органайзер|шкаф[-\s]?организ|супершкаф|кракен|смартбокс|кофе[-\s]?пойнт",
    "СУО / Face ID / Engy (ТП)":                  r"face\s?id|\bсуо\b|электронн\w*\s+очеред|\bengy\b|терминал\s+суо",
    "банкоматная зона (ТП)":                      r"банкомат|\bgrg\b|hyosung|cash\s?in|обрамлени\w*\s+банкомат|прибанкомат",
    "касса/депозитарий/инкассация (ТП)":          r"депозитар|кассов\w*\s+узел|операционн\w*\s+касс|инкассатор|кабина\s+клиента|преддепозитар",
    "укреплённость: броня/пулестойкость (ТП)":    r"бронестекл|бронедвер|бронеблок|пулестойк|взломостойк|устойчивост\w*\s+к\s+взлому|\bбр[-\s]?[235]\b|пуз[-\s]?бр|шлюз\s+инкассатор|передаточн\w*\s+лоток|защитн\w*\s+панел",
    "доводчики Dorma TS-83/TS-93 (ТП)":           r"ts[-\s]?93|ts[-\s]?83|dorma\s+ts",
    "мнемосхема (ТП)":                            r"мнемосхем",
    "декор ЛДСП Дуб Винченца H3157 / K006 (ТП)":  r"винченца|h\s?3157\b|k\s?006\b|дуб\s+урбан|urban\s+k\s?006|кроношпан",
    "трансформируемая перегородка / Dorma HSW":   r"dorma\s+hsw|hsw\s+easy|трансформируем\w*\s+перегород",
    "грязезащитный коврик СИТИТОП (ТП)":          r"сититоп|грязезащитн",
    "радиатор RIFAR (ТП)":                        r"\brifar\b",
    "вентустановка Globalvent/VTS/NED (ТП, тех. зона)": r"globalvent|\bvts\b|\bned\b",
    "сплит-система Haier / MDV (ТП, номинир.)":   r"\bhaier\b|\bmdv\b",
    "розетки Legrand Mozaik / DKC (ТП, номинир.)": r"mozaik|мозаик|legrand[\w\s-]{0,12}mozai",
    "двери Kapelli (ТП)":                         r"kapelli",
}
INFRA = {
    "акустический остров Ecophon Solo":           r"ecophon|акустическ\w*\s+остров",
    "ковролин Modulyss / Desso":                  r"modulyss|\bdesso\b",
    "ПВХ-плитка Vertigo/Gerflor/KBS (инфра)":     r"vertigo\s+loose|loose\s+lay|gerflor\s+creation|\bkbs\b[^\n]{0,20}dark",
    "плинтус Profilpas Metal Line 90 / DOLLKEN":  r"profilpas|metal\s+line\s+90|dollken|tle55",
    "фетровые обои BuzziSkin (инфра)":            r"buzziskin|фетров\w*\s+обо",
    "потолок Knauf П113 RAL 9010 (инфра)":        r"п\s?113[^\n]{0,25}9010|9010[^\n]{0,25}п\s?113|перфорированн\w*\s+акустическ\w*\s+потол|\bс3[-\s]?кр\b",
    "свет REFLECTA (Rota/Quark/Ray/Longa/…)":     r"reflecta|rota\s+(eco|d\d|egg|triangle|d\d{3})|quark\s+(air|in|it)|\bray\s+line\b|\bray\s+in\b|fat\s+comfort|\blonga\s+\d|agat\s+d\d|silent\s+acoustic|q45\s+track|grata\s+slim|mica\s+d190|nordlux|neon\s*flex",
    "цельностеклянная перегородка Pilkington/Versal": r"pilkington|\bversal\b|versal[-\s]?(arm|duplex)",
    "фурнитура Titan T-671 / N87D (инфра)":       r"titan[\s-]?t[\s-]?671|titan\s+n87d",
    "маркерное белое стекло RAL 9010 (инфра)":    r"маркерн\w*\s+(бел\w*\s+)?стекл|маркерное\s+бел",
    "спорт/керамогр.: Regupol/Coswick/Caesar":    r"regupol|everroll|coswick|caesar\s+inner|\bpeak\b",
    "керамогранит инфра (Про Дабл / Ломбардиа / Миллениум)": r"про\s+дабл|dd601200|ломбардиа|миллениум",
    "офисные зоны (фокус/кофе-поинт/ЛХЛК/…)":     r"фокус[-\s]?комнат|кофе[-\s]?поинт|\bлхлк\b|гардеробн|принтерн|спортзал|учебн\w*\s+класс|зона\s+отдыха|рабоч\w*\s+зона",
    "звукоизоляция Rw-45 dB (инфра, переговорн.)": r"rw[\s-]?45",
    "декор. штукатурка «под бетон» (инфра)":      r"под\s+бетон|антивандальн\w*\s+штукатур",
    "декор Egger Дуб Гладстоун / U999 (инфра)":   r"гладстоун|h3309|u999",
    "скрытые двери Hafele / Kubica (инфра)":      r"hafele|h[aä]fele|kubica",
    "дверь со звукоизол. заполнением (инфра)":    r"звукоизоляционн\w*\s+заполнени",
}
# «Слабые» признаки — только как подсказка, не как приговор.
KKO_WEAK   = {"цвет RAL 7047/3020 (палитра ТП)": r"ral\s*70?47|ral\s*30?20|ral\s*7021",
              "электроустановка Legrand/DKC/ABB/Schneider (палитра ТП)": r"\blegrand\b|\bdkc\b|schneider|\babb\b",
              "сантехприбор — в ТП-альбоме нормируется (SantiLine/GROHE/Ideal Standard), в смете аналог":
                  r"\bjika\b|jacob\s+delafon|santiline|\bgrohe\b|ideal\s+standard|\broca\b|cersanit|geberit|\bvitra\b|акватон|aquaton|am\.pm"}
INFRA_WEAK = {"цвет RAL 9010/9005 муар/9017 (палитра инфра)": r"ral\s*90?10|9005\s*муар|ral\s*90?17"}

MEP_SECT = ("Р5", "Р7")          # разделы, которых в альбомах почти нет (только оконечные устройства)


def hits(patterns, text):
    t = text.lower()
    return [name for name, rx in patterns.items() if re.search(rx, t)]


def norm(s):
    s = s.lower()
    s = re.sub(r"[0-9]+([.,][0-9]+)?", "#", s)          # числа/размеры -> #
    s = re.sub(r"[^\wа-яё ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb[SHEET]

    razd = pod = ""
    rows = []
    for r in ws.iter_rows(min_row=15, max_row=ws.max_row):
        A = r[0].value; B = r[1].value; G = r[6].value; P = r[15].value; Q = r[16].value
        gtext = (str(G).strip() if G is not None else "")
        if re.match(r"^Раздел\s*№", gtext):
            razd = gtext; pod = ""; continue
        if re.match(r"^\d+\.\d+", gtext):
            pod = gtext; continue
        if Q is None:
            continue
        q = str(Q).strip()
        if q == "" or re.fullmatch(r"[0-9]+([.,][0-9]+)?", q):
            continue
        rows.append(dict(row=r[0].row, razd=(A or razd), pod=pod,
                         name=gtext, note=(str(P).strip() if P else ""), q=q))

    # --- сигнатуры по каждой строке
    for d in rows:
        blob = d["name"] + " || " + d["note"]
        d["K"]  = hits(KKO,   blob)
        d["I"]  = hits(INFRA, blob)
        d["Kw"] = hits(KKO_WEAK,   blob)
        d["Iw"] = hits(INFRA_WEAK, blob)

    # --- поиск почти-дублей с разными тегами
    groups = {}
    for d in rows:
        groups.setdefault(norm(d["name"]), []).append(d)
    for g in groups.values():
        bases = {re.sub(r"\s*\?$", "", x["q"]) for x in g}
        if len(g) > 1 and len(bases) > 1:
            for x in g:
                x["dup"] = "да (в группе одинаковых работ теги разные: " + " / ".join(sorted(bases)) + ")"

    # --- вердикт
    for d in rows:
        base = re.sub(r"\s*\?$", "", d["q"])
        isq  = d["q"].endswith("?")
        sK, sI = bool(d["K"]), bool(d["I"])
        v = c = ""
        if base == "ККО, ИНФРАСТРУКТУРА":
            if sK and sI:      v, c = "OK", "признаки обоих альбомов — двойной тег обоснован"
            elif sK:           v, c = "УТОЧНИТЬ", "найдены признаки только ККО — двойной тег под вопросом"
            elif sI:           v, c = "УТОЧНИТЬ", "найдены признаки только ИНФРАСТРУКТУРА — двойной тег под вопросом"
            else:              v, c = "OK", "общая позиция без привязки к модели — применима к обоим"
        else:
            other = "ИНФРАСТРУКТУРА" if base == "ККО" else "ККО"
            mine   = sK if base == "ККО" else sI
            opp    = sI if base == "ККО" else sK
            mine_w = bool(d["Kw"]) if base == "ККО" else bool(d["Iw"])
            opp_w  = bool(d["Iw"]) if base == "ККО" else bool(d["Kw"])
            if opp and not mine:
                v = "ВЕРОЯТНО ОШИБКА"
                c = "по номенклатуре это " + other + " — совпало: " + "; ".join(d["I"] if base=="ККО" else d["K"])
            elif opp and mine:
                v = "УТОЧНИТЬ"
                c = "признаки обоих альбомов → возможно «ККО, ИНФРАСТРУКТУРА»"
            elif mine:
                v = "OK"
                c = "подтверждено: " + "; ".join(d["K"] if base=="ККО" else d["I"])
            elif mine_w and not opp_w:
                v = "СКОРЕЕ ВЕРНО"
                c = "точной модели в альбоме нет; косвенно за " + base + ": " + "; ".join(d["Kw"] if base=="ККО" else d["Iw"])
            else:
                v = "НЕ ПОДТВЕРЖДЕНО"
                if str(d["razd"]).startswith(MEP_SECT):
                    c = "инженерный раздел; в альбомах — только оконечные устройства, сверять не с чем"
                elif opp_w and not mine_w:
                    c = "точной модели нет; косвенные признаки указывают скорее на " + other + " (" + "; ".join(d["Iw"] if base=="ККО" else d["Kw"]) + ")"
                else:
                    c = "модели в справочниках нет — тег проставлен по смыслу раздела"
        if d.get("dup") and v == "OK":
            v = "УТОЧНИТЬ"
        d["verdict"], d["comment"] = v, c

    # --------------------------------------------------------------- запись
    out = openpyxl.Workbook()
    sh = out.active
    sh.title = "Сверка"
    head = ["Строка", "Раздел", "Подраздел", "Наименование работы", "Примечание",
            "Тег Q", "Знак «?»", "Вердикт", "Признаки ККО", "Признаки ИНФРАСТРУКТУРА",
            "Почти-дубль с другим тегом", "Комментарий"]
    sh.append(head)
    fills = {
        "OK":              PatternFill("solid", fgColor="E7F4E4"),
        "СКОРЕЕ ВЕРНО":    PatternFill("solid", fgColor="EFF7DF"),
        "УТОЧНИТЬ":        PatternFill("solid", fgColor="FFF4D6"),
        "ВЕРОЯТНО ОШИБКА": PatternFill("solid", fgColor="F8D7DA"),
        "НЕ ПОДТВЕРЖДЕНО": PatternFill("solid", fgColor="E2E3F0"),
    }
    for d in rows:
        sh.append([d["row"], d["razd"], d["pod"], d["name"], d["note"], d["q"],
                   "?" if d["q"].endswith("?") else "",
                   d["verdict"], "; ".join(d["K"]), "; ".join(d["I"]),
                   d.get("dup", ""), d["comment"]])
        for col in range(1, len(head) + 1):
            sh.cell(row=sh.max_row, column=col).fill = fills.get(d["verdict"], PatternFill())

    for c in sh[1]:
        c.font = Font(bold=True); c.alignment = Alignment(wrap_text=True, vertical="top")
    widths = [7, 8, 26, 60, 26, 20, 6, 17, 34, 34, 30, 60]
    for i, w in enumerate(widths, 1):
        sh.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    sh.freeze_panes = "A2"
    sh.auto_filter.ref = sh.dimensions

    # --- сводка
    sm = out.create_sheet("Сводка")
    from collections import Counter
    cv = Counter(d["verdict"] for d in rows)
    cbt = Counter((d["razd"], d["verdict"]) for d in rows)
    sm.append(["Вердикт", "Строк"])
    for k, n in cv.most_common():
        sm.append([k, n])
    sm.append([]); sm.append(["Раздел", "Вердикт", "Строк"])
    for (rz, vv), n in sorted(cbt.items()):
        sm.append([rz, vv, n])
    for c in sm[1]:
        c.font = Font(bold=True)
    sm.column_dimensions["A"].width = 20; sm.column_dimensions["B"].width = 20
    sm.column_dimensions["C"].width = 10

    out.save(OUT)
    print("saved:", OUT)
    for k, n in cv.most_common():
        print(f"  {k:18} {n}")


if __name__ == "__main__":
    main()
