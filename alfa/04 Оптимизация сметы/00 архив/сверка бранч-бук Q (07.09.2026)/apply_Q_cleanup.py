# -*- coding: utf-8 -*-
"""
Чистка столбца Q («Бранч-бук (ККО / Инфраструктура)») листа «Смета ЕР (общ)»
в мастер-файле  ТС ЕР исходник.xlsx  по результату сверки с бранч-буками.

Правило (вариант «явные из ББ + косвенно надёжные»):
  • вердикт OK / «СКОРЕЕ ВЕРНО»  -> тег остаётся, знак «?» снимается;
  • вердикт «НЕ ПОДТВЕРЖДЕНО»     -> ячейка Q очищается (стиль сохраняется).

Правка — на уровне zip/XML (переписывается только xl/worksheets/sheet1.xml),
поэтому проверка данных, картинки, стили и все прочие листы не затрагиваются.
Источник вердиктов — Сверка Q — построчно.xlsx.  Лог -> Q_cleanup_log.csv
"""
import csv, re, shutil, zipfile
from pathlib import Path
import openpyxl

BASE   = Path(__file__).resolve().parent
MASTER = BASE / "ТС ЕР исходник.xlsx"
SVERKA = BASE / "Сверка Q — построчно.xlsx"
LOG    = BASE / "Q_cleanup_log.csv"
SHEETXML = "xl/worksheets/sheet1.xml"   # лист «Смета ЕР (общ) » = rId1

# индексы sharedStrings (проверено в файле)
SI = {"ККО": 846, "ИНФРАСТРУКТУРА": 6176, "ККО ?": 6177,
      "ККО, ИНФРАСТРУКТУРА": 6178, "ИНФРАСТРУКТУРА ?": 6179}
STRIP = {"ККО ?": "ККО", "ИНФРАСТРУКТУРА ?": "ИНФРАСТРУКТУРА"}
KEEP = {"OK", "СКОРЕЕ ВЕРНО"}
DROP = {"НЕ ПОДТВЕРЖДЕНО"}


def plan():
    sh = openpyxl.load_workbook(SVERKA)["Сверка"]
    todo = []
    for row in sh.iter_rows(min_row=2, values_only=True):
        rnum, _rz, _pd, name, _nt, q, _qm, verdict = row[:8]
        if rnum is None:
            continue
        q, verdict = str(q), str(verdict)
        if verdict in DROP:
            todo.append((int(rnum), q, None, verdict, str(name)[:90]))
        elif verdict in KEEP and q in STRIP:
            todo.append((int(rnum), q, STRIP[q], verdict, str(name)[:90]))
        # OK/СКОРЕЕ ВЕРНО без «?» — без изменений
    return todo


def patch_xml(xml, todo):
    applied, problems = [], []
    for rnum, old, new, verdict, name in todo:
        pat = re.compile(r'<c r="Q%d"( s="(\d+)")?[^>]*?(?:/>|>.*?</c>)' % rnum, re.S)
        m = pat.search(xml)
        if not m:
            problems.append((rnum, "ячейка Q не найдена")); continue
        cur = m.group(0)
        vm = re.search(r'<v>(\d+)</v>', cur)
        cur_idx = int(vm.group(1)) if vm else None
        if cur_idx != SI.get(old):
            problems.append((rnum, f"в файле индекс {cur_idx}, ожидался {SI.get(old)} ({old!r})"))
            continue
        style = m.group(1) or ""
        if new is None:
            repl = f'<c r="Q{rnum}"{style}/>'
            act = "очищено"
        else:
            repl = f'<c r="Q{rnum}"{style} t="s"><v>{SI[new]}</v></c>'
            act = f'снят «?»  →  {new}'
        xml = xml[:m.start()] + repl + xml[m.end():]
        applied.append((rnum, old, "" if new is None else new, act, verdict, name))
    return xml, applied, problems


def rezip(src, dst, newxml):
    with zipfile.ZipFile(src) as zin, \
         zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename == SHEETXML:
                data = newxml.encode("utf-8")
            zi = zipfile.ZipInfo(it.filename, date_time=it.date_time)
            zi.compress_type = it.compress_type
            zi.external_attr = it.external_attr
            zi.internal_attr = it.internal_attr
            zi.create_system = it.create_system
            zout.writestr(zi, data)


def main():
    todo = plan()
    clr = sum(1 for t in todo if t[2] is None)
    print(f"К изменению: {len(todo)}  (очистить {clr}, снять «?» {len(todo)-clr})")

    with zipfile.ZipFile(MASTER) as z:
        xml = z.read(SHEETXML).decode("utf-8")
    newxml, applied, problems = patch_xml(xml, todo)

    if problems:
        print(f"ПРОБЛЕМЫ ({len(problems)}) — правка НЕ применена:")
        for rn, msg in problems[:20]:
            print(f"  r{rn}: {msg}")
        return

    bak = MASTER.with_suffix(".xlsx.bak_preQ")
    shutil.copy2(MASTER, bak)
    tmp = MASTER.with_suffix(".xlsx.tmp")
    rezip(MASTER, tmp, newxml)
    tmp.replace(MASTER)
    print(f"применено: {len(applied)}   (бэкап: {bak.name})")

    # верификация
    wb = openpyxl.load_workbook(MASTER)
    ws = wb["Смета ЕР (общ) "]
    bad = 0
    for rnum, old, new, act, verdict, name in applied:
        got = ws.cell(row=rnum, column=17).value
        exp = None if new == "" else new
        if (got or None) != exp:
            bad += 1; print(f"  ! r{rnum}: ожидалось {exp!r}, в файле {got!r}")
    left = sum(1 for r in ws.iter_rows(min_row=15, max_row=ws.max_row)
               if r[16].value and str(r[16].value).strip().endswith("?"))
    print(f"верификация: расхождений {bad};  строк со знаком «?» осталось: {left}")

    with open(LOG, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Строка", "Было", "Стало", "Действие", "Вердикт сверки", "Наименование"])
        for rec in applied:
            w.writerow(rec)
    print("лог:", LOG.name)


if __name__ == "__main__":
    main()
