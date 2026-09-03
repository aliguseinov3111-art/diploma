# -*- coding: utf-8 -*-
import openpyxl
wb = openpyxl.load_workbook("Форма КП ТЗ rev00.xlsx", data_only=False)
ws = wb["Форма КП"]
for r in range(85, 93):
    print(r, "C:", ws.cell(r,3).value)
    print("   D:", ws.cell(r,4).value)
