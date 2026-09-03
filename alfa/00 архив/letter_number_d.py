# -*- coding: utf-8 -*-
import openpyxl, re, string

wb = openpyxl.load_workbook("форма КП rev10.xlsx", data_only=False)
ws = wb["Форма КП"]

LEAD_RE = re.compile(r'^\d+(?:\.\d+)*\.?[а-яА-Яa-zA-Z]?\.?\s+')

current_block = None
groups = {}
order = []

for r in range(15, ws.max_row + 1):
    a = ws.cell(r, 1).value
    d = ws.cell(r, 4).value
    if a is not None and '.' in str(a):
        current_block = str(a).strip()
    if d is not None and str(d).strip() != '' and current_block is not None:
        groups.setdefault(current_block, [])
        if current_block not in order:
            order.append(current_block)
        groups[current_block].append(r)

max_group = max(len(v) for v in groups.values())
print("largest block size:", max_group)
assert max_group <= 26, "need multi-letter scheme"

changed = 0
for block in order:
    rows = groups[block]
    for i, r in enumerate(rows):
        letter = string.ascii_lowercase[i]
        old = ws.cell(r, 4).value
        stripped = LEAD_RE.sub('', old, count=1)
        new_val = f"{block}.{letter} {stripped}"
        if new_val != old:
            changed += 1
        ws.cell(r, 4).value = new_val

wb.save("форма КП rev10.xlsx")
print("blocks:", len(order))
print("rows numbered:", sum(len(v) for v in groups.values()))
print("rows changed:", changed)
